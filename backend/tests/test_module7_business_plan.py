"""
Module 7 business plan and pitch deck, verified offline.

The business plan asks Gemini (RIBIFModule7BusinessPlan.call_gemini) for CapEx,
OpEx, pricing and canvas text, then computes the P&L, NPV, IRR and payback
locally and writes an Excel workbook. The pitch deck sends a text brief to the
Presenton API over requests. Gemini and Presenton are faked here.
"""

import numpy_financial as npf
import pytest
import requests
from openpyxl import load_workbook

from utils.module_7.business_plan import BusinessPlanInput, RIBIFModule7BusinessPlan
from utils.module_7.pitch_deck import PitchDeckInput, RIBIFModule8PitchDeck

TCO_OUTPUT = {
    "project_name": "Profils PVC",
    "company_name": "Plastiques du Sahel",
    "total_tco": 180_000_000,
    "average_annual_tco": 36_000_000,
    "total_machines_tco": 120_000_000,
    "total_products_tco": 20_000_000,
    "total_raw_materials_tco": 40_000_000,
    "stage_1_informations_generales": {"informations_generales": {"description_projet": "Production de profils PVC pour menuiserie"}},
    "stage_3_produits_fabriques": {"produits_fabriques": [{"nom_produit": "Profil de fenetre"}]},
}

CAPEX = {"total_capex": 100_000_000, "equipment": {"items": ["Extrudeuse"], "total": 80_000_000}, "confidence": "high"}
OPEX = {"total_annual_opex": 40_000_000, "energy": {"description": "electricite", "annual_cost": 6_000_000}, "confidence": "high"}
PRICING = {
    "products": [{"name": "Profil de fenetre", "unit_price_dzd": 1200, "annual_volume": 60_000, "annual_revenue": 72_000_000, "growth_rate": 0.05}],
    "total_annual_revenue_year1": 72_000_000,
    "revenue_projections_5y": [72_000_000, 75_600_000, 79_380_000, 83_349_000, 87_516_450],
    "confidence": "high",
}


@pytest.fixture
def planner():
    return RIBIFModule7BusinessPlan(api_key="offline-test-key")


def _answer_by_prompt(prompt):
    if "total_capex" in prompt:
        return dict(CAPEX)
    if "total_annual_opex" in prompt:
        return dict(OPEX)
    if "revenue_projections_5y" in prompt:
        return dict(PRICING)
    return {}  # canvas text falls back to the built-in defaults


def test_business_plan_financials_follow_the_estimates(planner, monkeypatch, tmp_path):
    prompts = []

    def fake_call(prompt):
        prompts.append(prompt)
        return _answer_by_prompt(prompt)

    monkeypatch.setattr(planner, "call_gemini", fake_call)

    result = planner.generate_business_plan(BusinessPlanInput(**TCO_OUTPUT), output_dir=str(tmp_path))

    fin = result.financial_data
    assert fin["total_investment"] == 100_000_000
    assert fin["revenues"] == pytest.approx(PRICING["revenue_projections_5y"])
    assert fin["operating_costs"][0] == 40_000_000
    assert fin["cash_flows"][0] == -100_000_000
    assert fin["npv"] == pytest.approx(npf.npv(fin["discount_rate"], fin["cash_flows"]))
    assert result.total_revenue_5y_dzd == pytest.approx(sum(PRICING["revenue_projections_5y"]))
    assert any("Profils PVC" in p for p in prompts)

    workbook = load_workbook(result.filepath)
    assert len(workbook.sheetnames) == len(result.sheets) == 5


def test_business_plan_still_builds_when_gemini_is_down(planner, monkeypatch, tmp_path):
    monkeypatch.setattr(planner, "call_gemini", lambda prompt: {})

    result = planner.generate_business_plan(BusinessPlanInput(**TCO_OUTPUT), output_dir=str(tmp_path))

    assert result.capex_data["confidence"] == "low"
    assert result.size_bytes > 0


def test_business_plan_refuses_to_invent_tco(planner, tmp_path):
    with pytest.raises(ValueError, match="TCO"):
        planner.generate_business_plan(BusinessPlanInput(project_name="X", company_name="Y"), output_dir=str(tmp_path))


class PresentonReply:
    def __init__(self, status_code, payload=None, content=b""):
        self.status_code = status_code
        self._payload = payload or {}
        self.content = content
        self.text = str(payload)

    def json(self):
        return self._payload


def test_pitch_deck_sends_the_brief_and_saves_the_pdf(monkeypatch, tmp_path):
    monkeypatch.setenv("PRESENTON_API_KEY", "offline-test-key")
    sent = {}

    def fake_post(url, json=None, headers=None, timeout=None):
        sent.update(url=url, payload=json, headers=headers)
        return PresentonReply(200, {"presentation_id": "p1", "path": "/files/deck.pdf", "credits_consumed": 3})

    def fake_get(url, headers=None, timeout=None):
        sent["download"] = url
        return PresentonReply(200, content=b"%PDF-1.7 fake")

    monkeypatch.setattr(requests, "post", fake_post)
    monkeypatch.setattr(requests, "get", fake_get)

    deck = RIBIFModule8PitchDeck()
    result = deck.generate_pitch_deck(PitchDeckInput(**TCO_OUTPUT), output_dir=str(tmp_path))

    assert sent["url"].endswith("/api/v3/presentation/generate")
    assert sent["headers"]["Authorization"] == "Bearer offline-test-key"
    assert "Plastiques du Sahel" in sent["payload"]["content"]
    assert sent["payload"]["export_as"] == "pdf"
    assert sent["download"].endswith("/files/deck.pdf")
    assert result.pdf_size_bytes == len(b"%PDF-1.7 fake")
    assert result.credits_consumed == 3


def test_pitch_deck_reports_presenton_errors(monkeypatch, tmp_path):
    monkeypatch.setenv("PRESENTON_API_KEY", "offline-test-key")
    monkeypatch.setattr(requests, "post", lambda *a, **k: PresentonReply(402, {"detail": "no credits"}))

    with pytest.raises(Exception, match="402"):
        RIBIFModule8PitchDeck().generate_pitch_deck(PitchDeckInput(**TCO_OUTPUT), output_dir=str(tmp_path))


def test_pitch_deck_needs_an_api_key(monkeypatch, tmp_path):
    monkeypatch.delenv("PRESENTON_API_KEY", raising=False)

    with pytest.raises(ValueError, match="PRESENTON_API_KEY"):
        RIBIFModule8PitchDeck().generate_pitch_deck(PitchDeckInput(**TCO_OUTPUT), output_dir=str(tmp_path))
